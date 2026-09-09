from __future__ import annotations

import json
from pathlib import Path

import pytest
from openpyxl import Workbook

from g3_pipeline.builder import BuildError, build_dataset, verify_deterministic
from g3_pipeline.release_demo import release_demo_subset
from g3_pipeline.validator import validate_dataset

HEADERS = [
    "Dấu thời gian",
    "Địa chỉ email",
    "Tên thủ tục hành chính",
    "Lĩnh vực",
    "Hình thức nộp",
    "Thành phần hồ sơ",
    "Thời gian giải quyết",
    "Lệ phí",
    "Địa điểm tiếp nhận hồ sơ trực tiếp (nếu có)",
    "Ghi chú",
    "File biểu mẫu",
    "Trạng thái",
]

BUSINESS_HEADERS = HEADERS[2:11]


def create_workbook(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Dữ liệu mẫu"
    sheet.append(HEADERS)
    sheet.append(
        [
            "07/10/2025 08:00",
            "owner@example.test",
            "Thủ tục Mặt Trăng",
            "Hộ tịch",
            "Cả hai",
            "- Giấy tờ Alpha\n- Giấy tờ Beta",
            "3 ngày làm việc",
            "Không thu phí",
            "Bộ phận Một cửa",
            None,
            None,
            "Xong",
        ]
    )
    sheet.append(
        [
            None,
            "second@example.test",
            "Thủ tục Sao Mai",
            "Giáo dục và đào tạo",
            "Trực tuyến",
            "Giấy tờ Gamma",
            "5 ngày làm việc",
            "10.000 đồng",
            None,
            None,
            None,
            "Xong",
        ]
    )
    sheet.append(
        [
            None,
            "third@example.test",
            "Thủ tục Trùng",
            "Văn hóa - Xã hội",
            "Trực tiếp",
            "Giấy tờ Một",
            "2 ngày làm việc",
            "Không",
            "Điểm tiếp nhận A",
            None,
            None,
            "Xong",
        ]
    )
    sheet.append(
        [
            None,
            "fourth@example.test",
            "Thủ tục Trùng",
            "Văn hóa - Xã hội",
            "Trực tiếp",
            "Giấy tờ Hai",
            "4 ngày làm việc",
            "Không",
            "Điểm tiếp nhận B",
            None,
            None,
            "Xong",
        ]
    )
    sheet.append([None] * len(HEADERS))
    workbook.save(path)


def create_nine_column_workbook(path: Path) -> None:
    """Build a supported nine-business-column fixture without form metadata."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Dữ liệu mẫu 9 cột"
    sheet.append(BUSINESS_HEADERS)
    sheet.append(
        [
            "Thủ tục Mặt Trăng",
            "Hộ tịch",
            "Cả hai",
            "- Giấy tờ Alpha\n- Giấy tờ Beta",
            "3 ngày làm việc",
            "Không thu phí",
            "Bộ phận Một cửa",
            None,
            None,
        ]
    )
    workbook.save(path)


def schema_path() -> Path:
    return Path(__file__).parents[3] / "docs/giai-doan-3/schema/data-contract-v1.schema.json"


def load_jsonl(path: Path) -> list[dict]:
    return [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines()]


def test_build_is_schema_valid_private_and_deterministic(tmp_path: Path) -> None:
    workbook = tmp_path / "synthetic.xlsx"
    create_workbook(workbook)
    result = build_dataset(workbook, tmp_path / "private")

    report = validate_dataset(
        result.dataset_dir,
        schema_path(),
        mode="preflight",
        expected_source_count=4,
        workbook_path=workbook,
    )

    assert report["status"] == "PASS", report
    assert report["counts"] == {
        "sources": 4,
        "procedures": 3,
        "fragments": 19,
        "conflicts": 1,
        "questions": 9,
    }
    assert report["warnings"]["PROCEDURE_REVIEW_PENDING"] == 3
    assert report["warnings"]["GOLD_REVIEW_PENDING"] == 9
    assert "example.test" not in "".join(
        path.read_text(encoding="utf-8") for path in result.dataset_dir.rglob("*") if path.is_file()
    )
    assert verify_deterministic(workbook)


def test_review_is_never_approved_automatically(tmp_path: Path) -> None:
    workbook = tmp_path / "synthetic.xlsx"
    create_workbook(workbook)
    result = build_dataset(workbook, tmp_path / "private")
    procedures = load_jsonl(result.dataset_dir / "normalized/procedures.jsonl")
    statuses = {row["review"]["status"] for row in procedures}
    assert statuses == {"PENDING_REVIEW", "QUARANTINED"}


def test_build_accepts_nine_recognized_business_columns_without_form_metadata(
    tmp_path: Path,
) -> None:
    workbook = tmp_path / "synthetic-nine-columns.xlsx"
    create_nine_column_workbook(workbook)

    result = build_dataset(workbook, tmp_path / "private")
    report = validate_dataset(
        result.dataset_dir,
        schema_path(),
        mode="preflight",
        expected_source_count=1,
        workbook_path=workbook,
    )
    sources = load_jsonl(result.dataset_dir / "source_manifest.jsonl")

    assert report["status"] == "PASS", report
    assert report["counts"]["sources"] == 1
    assert sources[0]["collected_at"] is None


def test_unknown_nonempty_column_is_rejected_instead_of_silently_dropped(
    tmp_path: Path,
) -> None:
    workbook_path = tmp_path / "synthetic-unknown-column.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append([*BUSINESS_HEADERS, "Cột nghiệp vụ chưa có mapping"])
    sheet.append(
        [
            "Thủ tục Mặt Trăng",
            "Hộ tịch",
            "Trực tiếp",
            "Giấy tờ Alpha",
            "3 ngày làm việc",
            "Không thu phí",
            "Bộ phận Một cửa",
            None,
            None,
            "Dữ kiện không được phép bỏ qua",
        ]
    )
    workbook.save(workbook_path)

    with pytest.raises(BuildError, match="UNSUPPORTED_COLUMNS_WITH_DATA_AT_INDEX:10"):
        build_dataset(workbook_path, tmp_path / "private")


def test_validator_rejects_pii_in_gold(tmp_path: Path) -> None:
    workbook = tmp_path / "synthetic.xlsx"
    create_workbook(workbook)
    result = build_dataset(workbook, tmp_path / "private")
    questions_path = result.dataset_dir / "gold/questions.jsonl"
    questions = load_jsonl(questions_path)
    questions[0]["question"] += " owner@example.test"
    questions_path.write_text(
        "".join(json.dumps(row, ensure_ascii=False) + "\n" for row in questions),
        encoding="utf-8",
        newline="\n",
    )

    report = validate_dataset(result.dataset_dir, schema_path())

    assert report["status"] == "FAIL"
    assert report["errors"]["SENSITIVE_PATTERN_EMAIL"] == 1


def test_publish_requires_human_review(tmp_path: Path) -> None:
    workbook = tmp_path / "synthetic.xlsx"
    create_workbook(workbook)
    result = build_dataset(workbook, tmp_path / "private")

    report = validate_dataset(result.dataset_dir, schema_path(), mode="publish")

    assert report["status"] == "FAIL"
    assert report["errors"]["PROCEDURE_REVIEW_PENDING"] == 3
    assert report["errors"]["GOLD_REVIEW_PENDING"] == 9


def test_demo_release_is_reviewed_publishable_and_split_by_procedure(tmp_path: Path) -> None:
    workbook = tmp_path / "synthetic.xlsx"
    create_workbook(workbook)
    private_root = tmp_path / "company_private"
    source = build_dataset(workbook, private_root / "source")

    release = release_demo_subset(
        source.dataset_dir,
        private_root / "releases",
        schema_path(),
        reviewer_role="PROJECT_OWNER_DELEGATED_REVIEW",
        reviewed_at="2026-09-09T19:00:00+07:00",
        demo_size=2,
    )
    report = validate_dataset(
        release.dataset_dir,
        schema_path(),
        mode="publish",
        expected_source_count=2,
    )
    procedures = load_jsonl(release.dataset_dir / "normalized/procedures.jsonl")
    questions = load_jsonl(release.dataset_dir / "gold/questions.jsonl")
    conflicts = load_jsonl(release.dataset_dir / "quarantine/conflicts.jsonl")
    split_by_procedure: dict[str, set[str]] = {}
    for question in questions:
        if question["procedure_id"] is not None:
            split_by_procedure.setdefault(question["procedure_id"], set()).add(question["split"])

    assert report["status"] == "PASS", report
    assert len(procedures) == 2
    assert conflicts == []
    assert all(row["review"]["status"] == "APPROVED_FOR_DEMO" for row in procedures)
    assert all(len(splits) == 1 for splits in split_by_procedure.values())
    assert {next(iter(splits)) for splits in split_by_procedure.values()} == {"DEV", "TEST"}
    assert release.report["test_lock_sha256"]
